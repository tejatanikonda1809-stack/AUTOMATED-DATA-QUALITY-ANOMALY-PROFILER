"""Missing Values Audit Log Generator (openpyxl)."""
import io
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from backend.utils.df_utils import get_missing_mask, is_identifier_col
from backend.comparison.row_matcher import detect_identifier_column


def build_missing_values_workbook(orig_df: pd.DataFrame, final_df: pd.DataFrame, history: list = None, dataset_name: str = "Dataset") -> openpyxl.Workbook:
    """Generate an Excel workbook auditing every cell originally missing."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Missing Values Audit Log"
    ws.views.sheetView[0].showGridLines = True

    history = history or []

    # Detect identifier column for row tracking
    id_col = detect_identifier_column(orig_df)

    # 1. Title Block
    title_font = Font(name="Segoe UI", size=16, bold=True, color="1E293B")
    sub_font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
    
    ws["A1"] = f"Missing Values Audit Log — {dataset_name.upper()}"
    ws["A1"].font = title_font
    ws["A2"] = f"Comprehensive cell-level audit of all originally-missing values and their resolution."
    ws["A2"].font = sub_font

    # 2. Table Headers
    headers = [
        "Column Name",
        "Row Identifier",
        "Original Row #",
        "Original State",
        "Resolution Method",
        "Final Value",
        "Resolution Status"
    ]
    
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Deep blue
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    start_row = 4
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border
    ws.row_dimensions[start_row].height = 26

    # 3. Analyze History Actions
    fill_actions = {}  # col -> list of actions
    dropped_cols = set()
    deleted_row_indices = set()
    cell_edits = {}  # (r, c) -> new_val

    for entry in history:
        act = entry.get("action", "")
        col = entry.get("column", "")
        if act == "FILL_MISSING":
            method = entry.get("method", "imputation")
            fill_val = entry.get("fill_value")
            desc = entry.get("description", f"Filled using {method}")
            fill_actions.setdefault(col, []).append({
                "method": method,
                "fill_value": fill_val,
                "description": desc,
            })
        elif act in ("DELETE_COLUMN", "DROP_COLUMN"):
            dropped_cols.add(col)
        elif act in ("DELETE_ROWS", "DROP_ROWS"):
            for idx in entry.get("indices", []):
                try:
                    deleted_row_indices.add(int(idx))
                except (ValueError, TypeError):
                    pass
        elif act == "EDIT_CELL":
            r = entry.get("row_index")
            c = entry.get("column")
            if r is not None and c:
                cell_edits[(int(r), str(c))] = entry.get("new_value")

    # 4. Identify all originally missing cells
    missing_mask = get_missing_mask(orig_df)
    missing_cells = []
    for r_idx in orig_df.index:
        for c in orig_df.columns:
            if bool(missing_mask.loc[r_idx, c]):
                missing_cells.append((r_idx, c))

    # Pre-index final_df by id_col if possible
    final_by_id = None
    if id_col and id_col in final_df.columns and id_col in orig_df.columns:
        try:
            final_by_id = final_df.set_index(id_col)
        except Exception:
            final_by_id = None

    row_font = Font(name="Segoe UI", size=10, color="0F172A")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    regular_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    status_resolved_font = Font(name="Segoe UI", size=10, bold=True, color="166534")  # Green
    status_dropped_font = Font(name="Segoe UI", size=10, bold=True, color="854D0E")   # Amber
    status_unresolved_font = Font(name="Segoe UI", size=10, bold=True, color="991B1B") # Red

    curr_row = start_row + 1

    if not missing_cells:
        ws.cell(row=curr_row, column=1, value="No missing values were detected in the original dataset.")
        ws.cell(row=curr_row, column=1).font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=len(headers))
        curr_row += 1
    else:
        for orig_r_idx, col in missing_cells:
            # Row Identifier
            if id_col and id_col in orig_df.columns:
                id_val = str(orig_df.loc[orig_r_idx, id_col])
                row_ident = f"{id_val} (Row {orig_r_idx})"
            else:
                row_ident = f"Row {orig_r_idx}"

            # Determine resolution
            final_val = None
            resolution_method = "Unresolved"
            status = "Unresolved"

            # Check if column was dropped
            if col not in final_df.columns or col in dropped_cols:
                resolution_method = "Column Dropped"
                final_val = "[Column Removed]"
                status = "Removed via Column Drop"

            # Check if row was deleted
            elif orig_r_idx in deleted_row_indices:
                resolution_method = "Row Deleted"
                final_val = "[Row Deleted]"
                status = "Removed via Row Deletion"

            else:
                # Check if row exists in final_df
                row_found = False
                cur_val = None

                if final_by_id is not None and id_col in orig_df.columns:
                    id_val = orig_df.loc[orig_r_idx, id_col]
                    if id_val in final_by_id.index:
                        row_found = True
                        val_series = final_by_id.loc[id_val, col]
                        cur_val = (val_series.iloc[0] if len(val_series) > 0 else None) if isinstance(val_series, pd.Series) else val_series
                
                if not row_found and orig_r_idx in final_df.index:
                    row_found = True
                    cur_val = final_df.loc[orig_r_idx, col]

                if not row_found:
                    # Row no longer present
                    resolution_method = "Row Removed"
                    final_val = "[Row Removed]"
                    status = "Removed via Row Deletion"
                else:
                    # Check if cell is still missing
                    is_still_missing = bool(get_missing_mask(pd.Series([cur_val]))[0])
                    if is_still_missing:
                        resolution_method = "Unresolved (Still Missing)"
                        final_val = "[Missing / Null]"
                        status = "Unresolved"
                    else:
                        final_val = str(cur_val)
                        status = "Resolved"
                        if (orig_r_idx, col) in cell_edits:
                            resolution_method = f"Manual Cell Edit ('{cell_edits[(orig_r_idx, col)]}')"
                        elif col in fill_actions:
                            action_info = fill_actions[col][-1]
                            resolution_method = action_info.get("description") or f"Imputed ({action_info.get('method')})"
                        else:
                            resolution_method = "Imputed / Cleaned"

            row_data = [
                str(col),
                row_ident,
                str(orig_r_idx),
                "Missing (Null / Empty)",
                resolution_method,
                str(final_val) if final_val is not None else "",
                status
            ]

            fill = alt_fill if (curr_row % 2 == 0) else regular_fill
            for c_idx, val in enumerate(row_data, start=1):
                c = ws.cell(row=curr_row, column=c_idx, value=val)
                c.fill = fill
                c.font = row_font
                c.border = thin_border
                c.alignment = Alignment(vertical="center")

                # Style status column
                if c_idx == 7:
                    if "Resolved" in status:
                        c.font = status_resolved_font
                    elif "Removed" in status:
                        c.font = status_dropped_font
                    else:
                        c.font = status_unresolved_font

            ws.row_dimensions[curr_row].height = 20
            curr_row += 1

    # Auto-adjust column widths
    for col_cells in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col_cells[0].column)
        for c in col_cells:
            val_str = str(c.value or "")
            if len(val_str) > max_len and c.row >= 4:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    return wb


def export_missing_values_log_bytes(orig_df: pd.DataFrame, final_df: pd.DataFrame, history: list = None, dataset_name: str = "Dataset") -> bytes:
    wb = build_missing_values_workbook(orig_df, final_df, history, dataset_name)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
